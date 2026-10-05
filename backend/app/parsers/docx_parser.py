import io
import docx

def extract_text_from_docx(file_input):
    """
    Extracts text from a DOCX file path or file-like byte stream.
    Returns (text, paragraph_count, metadata, error).
    """
    try:
        if isinstance(file_input, (bytes, bytearray)):
            stream = io.BytesIO(file_input)
            doc = docx.Document(stream)
        elif isinstance(file_input, str):
            doc = docx.Document(file_input)
        else:
            doc = docx.Document(file_input)
            
        text_parts = []
        
        # Paragraphs
        for p in doc.paragraphs:
            txt = p.text.strip()
            if txt:
                text_parts.append(txt)
                
        # Tables
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    # Deduplicate repeated text from merged cells
                    unique_row_texts = []
                    for rt in row_texts:
                        if not unique_row_texts or rt != unique_row_texts[-1]:
                            unique_row_texts.append(rt)
                    if unique_row_texts:
                        text_parts.append(" | ".join(unique_row_texts))
                        
        full_text = "\n".join(text_parts).strip()
        
        metadata = {}
        if doc.core_properties:
            if doc.core_properties.title:
                metadata['title'] = doc.core_properties.title
            if doc.core_properties.author:
                metadata['author'] = doc.core_properties.author
                
        if not full_text:
            return "", 0, metadata, "Unable to extract text. The document appears to be empty."
            
        return full_text, len(doc.paragraphs), metadata, None
        
    except Exception as e:
        return "", 0, {}, f"Error reading DOCX: {str(e)}"
