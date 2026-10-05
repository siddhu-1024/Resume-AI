import os
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import docx

def create_sample_pdf(file_path, name, email, phone, location, summary, skills, experience_entries, edu, projects=None, certs=None):
    doc = SimpleDocTemplate(file_path, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle('Title', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=18, leading=22)
    contact_style = ParagraphStyle('Contact', parent=styles['Normal'], fontName='Helvetica', fontSize=9, leading=13, textColor='#475569')
    head_style = ParagraphStyle('Head', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=12, leading=16, spaceBefore=10, spaceAfter=4, textColor='#1E293B')
    body_style = ParagraphStyle('Body', parent=styles['Normal'], fontName='Helvetica', fontSize=9, leading=13, textColor='#334155')
    
    elements = []
    
    # Header
    elements.append(Paragraph(f"<b>{name}</b>", title_style))
    elements.append(Paragraph(f"{location} | {email} | {phone}", contact_style))
    elements.append(Spacer(1, 10))
    
    # Summary
    elements.append(Paragraph("PROFESSIONAL SUMMARY", head_style))
    elements.append(Paragraph(summary, body_style))
    elements.append(Spacer(1, 8))
    
    # Skills
    elements.append(Paragraph("TECHNICAL SKILLS", head_style))
    elements.append(Paragraph(f"<b>Proficiencies:</b> {', '.join(skills)}", body_style))
    elements.append(Spacer(1, 8))
    
    # Experience
    elements.append(Paragraph("WORK EXPERIENCE", head_style))
    for exp in experience_entries:
        elements.append(Paragraph(f"<b>{exp['title']}</b> &bull; {exp['company']} ({exp['dates']})", body_style))
        for bullet in exp['bullets']:
            elements.append(Paragraph(f"&bull; {bullet}", body_style))
        elements.append(Spacer(1, 4))
    elements.append(Spacer(1, 4))
    
    # Education
    elements.append(Paragraph("EDUCATION", head_style))
    elements.append(Paragraph(f"<b>{edu['degree']}</b> &bull; {edu['school']} ({edu['year']})", body_style))
    elements.append(Spacer(1, 8))
    
    if projects:
        elements.append(Paragraph("KEY PROJECTS", head_style))
        for p in projects:
            elements.append(Paragraph(f"<b>{p['title']}</b> - {p['desc']}", body_style))
        elements.append(Spacer(1, 8))
        
    if certs:
        elements.append(Paragraph("CERTIFICATIONS", head_style))
        for c in certs:
            elements.append(Paragraph(f"&bull; {c}", body_style))
            
    doc.build(elements)

def create_sample_docx(file_path, name, email, phone, location, summary, skills, experience_entries, edu, projects=None, certs=None):
    doc = docx.Document()
    
    # Header
    doc.add_heading(name, level=1)
    p_head = doc.add_paragraph(f"{location} | {email} | {phone}")
    
    # Summary
    doc.add_heading("Professional Summary", level=2)
    doc.add_paragraph(summary)
    
    # Skills
    doc.add_heading("Technical Skills", level=2)
    doc.add_paragraph(", ".join(skills))
    
    # Experience
    doc.add_heading("Work Experience", level=2)
    for exp in experience_entries:
        p = doc.add_paragraph()
        r = p.add_run(f"{exp['title']} at {exp['company']} ({exp['dates']})\n")
        r.bold = True
        for b in exp['bullets']:
            doc.add_paragraph(b, style='List Bullet')
            
    # Education
    doc.add_heading("Education", level=2)
    doc.add_paragraph(f"{edu['degree']}, {edu['school']} ({edu['year']})")
    
    if projects:
        doc.add_heading("Key Projects", level=2)
        for pr in projects:
            doc.add_paragraph(f"{pr['title']}: {pr['desc']}", style='List Bullet')
            
    if certs:
        doc.add_heading("Certifications", level=2)
        for c in certs:
            doc.add_paragraph(c, style='List Bullet')
            
    doc.save(file_path)

def generate_all_samples(output_dir="sample_resumes"):
    os.makedirs(output_dir, exist_ok=True)
    
    # 1. Strong Full-Stack Python PDF
    create_sample_pdf(
        file_path=os.path.join(output_dir, "sample_strong_fullstack.pdf"),
        name="Elena Rostova",
        email="elena.rostova@gmail.com",
        phone="+1 (415) 781-9920",
        location="San Francisco, CA",
        summary="Senior Full-Stack Engineer with 6+ years of experience building scalable cloud microservices in Python (FastAPI/Flask), robust PostgreSQL database architectures, and reactive React applications.",
        skills=["Python", "FastAPI", "Flask", "Django", "PostgreSQL", "Redis", "Docker", "Kubernetes", "AWS", "REST APIs", "React", "TypeScript", "Pytest", "CI/CD", "Git", "System Design"],
        experience_entries=[
            {
                "title": "Lead Software Engineer",
                "company": "Nexus Cloud Systems",
                "dates": "2021 - Present",
                "bullets": [
                    "Designed and maintained high-throughput REST APIs in Python FastAPI backed by PostgreSQL and Redis.",
                    "Orchestrated Docker containers and deployed automated Kubernetes pipelines on AWS.",
                    "Mentored 5 junior developers and conducted architecture reviews."
                ]
            },
            {
                "title": "Backend Developer",
                "company": "Apex Fintech",
                "dates": "2018 - 2021",
                "bullets": [
                    "Developed Python Flask microservices and integrated third-party banking REST APIs.",
                    "Implemented comprehensive unit test suites with pytest achieving 95% code coverage."
                ]
            }
        ],
        edu={"degree": "Bachelor of Science in Computer Science", "school": "University of California, Berkeley", "year": "2018"},
        projects=[
            {"title": "Real-time Order Matching Engine", "desc": "High-frequency transaction service built in Python, Redis, and React."}
        ],
        certs=["AWS Certified Solutions Architect", "Certified Kubernetes Application Developer"]
    )
    
    # 2. Frontend React DOCX
    create_sample_docx(
        file_path=os.path.join(output_dir, "sample_frontend_developer.docx"),
        name="David Kim",
        email="david.kim.design@outlook.com",
        phone="+1 (206) 910-3482",
        location="Seattle, WA",
        summary="Senior Frontend Engineer with 4 years of expertise creating web applications with React, TypeScript, Next.js, and modern CSS frameworks. Strong advocate for component-driven UI and automated frontend testing.",
        skills=["React", "TypeScript", "JavaScript", "Next.js", "Tailwind CSS", "Redux", "HTML5", "CSS3", "Jest", "Cypress", "GraphQL", "REST APIs", "Git", "Agile"],
        experience_entries=[
            {
                "title": "Senior Frontend Developer",
                "company": "Verve Web Studios",
                "dates": "2022 - Present",
                "bullets": [
                    "Architected modern component library using React, Next.js, and Tailwind CSS.",
                    "Reduced core bundle load times by 38% through code splitting and asset optimization.",
                    "Wrote automated UI unit tests using Jest and Cypress."
                ]
            },
            {
                "title": "UI Engineer",
                "company": "PixelCraft Apps",
                "dates": "2020 - 2022",
                "bullets": [
                    "Built reactive interfaces with React and integrated REST APIs.",
                    "Collaborated with UX designers to implement responsive designs."
                ]
            }
        ],
        edu={"degree": "Bachelor of Science in Software Engineering", "school": "University of Washington", "year": "2020"},
        projects=[
            {"title": "SaaS Design System", "desc": "Accessible UI component library built in React and TypeScript."}
        ]
    )

    # 3. Junior Python PDF
    create_sample_pdf(
        file_path=os.path.join(output_dir, "sample_junior_python.pdf"),
        name="Lucas Silva",
        email="lucas.silva98@yahoo.com",
        phone="+1 (617) 321-4902",
        location="Boston, MA",
        summary="Enthusiastic junior developer with foundational coursework and internship experience in Python, SQLite, HTML/CSS, and basic web scripting.",
        skills=["Python", "JavaScript", "HTML", "CSS", "SQLite", "Git", "Problem Solving", "Teamwork"],
        experience_entries=[
            {
                "title": "Software Engineering Intern",
                "company": "InnoSoft Labs",
                "dates": "2023 - 2024",
                "bullets": [
                    "Assisted with Python automation scripts and basic bug triage.",
                    "Developed static web pages using HTML, CSS, and vanilla JavaScript."
                ]
            }
        ],
        edu={"degree": "Bachelor of Science in Computer Science", "school": "Boston University", "year": "2023"}
    )

    # 4. Data Engineer DOCX
    create_sample_docx(
        file_path=os.path.join(output_dir, "sample_data_engineer.docx"),
        name="Priya Patel",
        email="priya.patel.data@gmail.com",
        phone="+1 (408) 555-7839",
        location="San Jose, CA",
        summary="Data & ML Engineer with 5 years experience building data pipelines with Python, Apache Spark, Kafka, SQL, and deploying machine learning models.",
        skills=["Python", "SQL", "PostgreSQL", "Spark", "Kafka", "Docker", "AWS", "Airflow", "Scikit-Learn", "Pandas", "NumPy", "REST APIs", "Git"],
        experience_entries=[
            {
                "title": "Senior Data Engineer",
                "company": "DataPulse Corp",
                "dates": "2021 - Present",
                "bullets": [
                    "Built streaming data pipelines in Python and Apache Spark processing 10TB+ daily data.",
                    "Orchestrated DAG workflows with Apache Airflow and managed PostgreSQL databases."
                ]
            }
        ],
        edu={"degree": "Master of Science in Data Analytics", "school": "San Jose State University", "year": "2019"}
    )

    # 5. Unrelated Sales PDF
    create_sample_pdf(
        file_path=os.path.join(output_dir, "sample_unrelated_candidate.pdf"),
        name="Arthur Pendelton",
        email="arthur.pendelton@gmail.com",
        phone="+1 (702) 441-2900",
        location="Las Vegas, NV",
        summary="Experienced Hospitality & Retail Operations Manager with 8+ years leading customer service teams, inventory planning, and guest relations.",
        skills=["Customer Service", "Sales Management", "Inventory Control", "Staff Scheduling", "Event Planning", "Negotiation", "Leadership"],
        experience_entries=[
            {
                "title": "General Manager",
                "company": "Grand Vista Resorts",
                "dates": "2018 - Present",
                "bullets": [
                    "Managed 40+ hospitality staff and oversaw $5M annual retail operations.",
                    "Maintained 98% customer satisfaction score across resort properties."
                ]
            }
        ],
        edu={"degree": "Bachelor of Arts in Hospitality Management", "school": "UNLV", "year": "2016"}
    )

if __name__ == '__main__':
    generate_all_samples()
    print("Sample resumes generated successfully.")
